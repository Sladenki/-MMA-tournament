from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from mma_secretary.core.normalize import normalize_name, normalize_org, normalize_rank, parse_weight
from mma_secretary.services.engine import TournamentService

HEADERS = ["№", "жребий", "контрольный", "Фамилия, имя", "Город, организация", "Разряд", "Год", "Тренер", "Вес"]


def import_xlsx(svc: TournamentService, data: bytes | Path) -> dict:
    if isinstance(data, Path):
        wb = load_workbook(data, data_only=True)
    else:
        wb = load_workbook(BytesIO(data), data_only=True)
    ws = wb[wb.sheetnames[0]]
    added, skipped = [], []
    start = 1
    # detect header row
    for r in range(1, min(15, ws.max_row or 1) + 1):
        vals = [str(ws.cell(r, c).value or "").lower() for c in range(1, 10)]
        if any("фамил" in v or "фио" in v for v in vals):
            start = r + 1
            break
        if ws.cell(r, 2).value == 1 or ws.cell(r, 1).value == 1:
            start = r
            break
    empty = 0
    for r in range(start, (ws.max_row or start) + 1):
        # try B-J layout (old file) first: B=n C=draw D=ctrl E=name F=org G=rank H=year I=coach J=weight
        name = ws.cell(r, 5).value or ws.cell(r, 4).value
        if not name:
            # A-I layout
            name = ws.cell(r, 4).value
        if not name:
            empty += 1
            if empty >= 2:
                break
            continue
        empty = 0
        # old file columns B-J if column E looks like a name
        if ws.cell(r, 5).value and isinstance(ws.cell(r, 5).value, str):
            rec = {
                "seq": ws.cell(r, 2).value,
                "draw_number": ws.cell(r, 3).value,
                "name": ws.cell(r, 5).value,
                "organization": ws.cell(r, 6).value,
                "rank": ws.cell(r, 7).value,
                "birth_year": ws.cell(r, 8).value,
                "coach": ws.cell(r, 9).value,
                "weight": ws.cell(r, 10).value,
            }
        else:
            rec = {
                "seq": ws.cell(r, 1).value,
                "draw_number": ws.cell(r, 2).value,
                "name": ws.cell(r, 4).value,
                "organization": ws.cell(r, 5).value,
                "rank": ws.cell(r, 6).value,
                "birth_year": ws.cell(r, 7).value,
                "coach": ws.cell(r, 8).value,
                "weight": ws.cell(r, 9).value,
            }
        rec["name"] = normalize_name(rec["name"])
        if not rec["name"]:
            skipped.append({"row": r, "reason": "нет ФИО"})
            continue
        rec["organization"] = normalize_org(rec["organization"])
        rec["rank"] = normalize_rank(rec["rank"])
        rec["weight"] = parse_weight(rec["weight"])
        try:
            rec["birth_year"] = int(rec["birth_year"]) if rec["birth_year"] not in (None, "") else None
        except (TypeError, ValueError):
            rec["birth_year"] = None
        try:
            rec["draw_number"] = int(rec["draw_number"]) if rec["draw_number"] not in (None, "") else None
        except (TypeError, ValueError):
            rec["draw_number"] = None
        added.append(svc.add_participant(rec))
    return {"added": len(added), "skipped": skipped, "participants": added}


def export_participants_xlsx(svc: TournamentService) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Участники"
    ws.append(HEADERS)
    for p in svc.list_participants():
        ws.append(
            [
                p["seq"],
                p["draw_number"],
                "",
                p["name"],
                p["organization"],
                p["rank"],
                p["birth_year"],
                p["coach"],
                p["weight"],
            ]
        )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_report_xlsx(svc: TournamentService) -> bytes:
    wb = Workbook()
    # mandate
    man = svc.mandate()
    ws = wb.active
    ws.title = "Мандат"
    t = svc.get_tournament()
    ws.append([t.get("name"), t.get("kind")])
    ws.append([t.get("date"), t.get("city")])
    ws.append(["Организация", *man["weights"], "Итого", "Представитель"])
    for row in man["rows"]:
        ws.append([row["organization"], *[row["cells"].get(w, 0) for w in man["weights"]], row["total"], row["representative"]])
    ws.append(["Итого", *[man["col_totals"].get(w, 0) for w in man["weights"]], man["grand_total"], ""])
    ws.append([])
    ws.append(["Звания"])
    for k, v in man["titles"].items():
        ws.append([k, v])

    ws2 = wb.create_sheet("Команды")
    teams = svc.team_report()
    ws2.append(["Место", "Организация", "Очки", "Представитель"])
    for s in teams:
        ws2.append([s["place"], s["organization"], s["points"], s["representative"]])

    ws3 = wb.create_sheet("Личное")
    ws3.append(["Категория", "Место", "ФИО", "Год", "Разряд", "Организация", "Тренер", "Очки"])
    for r in svc.placements_report():
        cat = f"{r['age_label']} {r['division_code']} {r['weight_label']}"
        ws3.append([cat, r["place"], r["name"], r["birth_year"], r["rank"], r["organization"], r["coach"], r["points"]])

    ws4 = wb.create_sheet("Призёры")
    ws4.append(["Категория", "Место", "ФИО", "Организация", "Тренер"])
    for r in svc.medalists():
        cat = f"{r['age_label']} {r['division_code']} {r['weight_label']}"
        ws4.append([cat, r["place"], r["name"], r["organization"], r["coach"]])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
